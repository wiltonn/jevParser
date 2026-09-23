"""Change-event golden fixture -- the differ's actual output contract.

Every case here is one a naive diff gets wrong, and every one is resolved
deterministically: no model is consulted to produce any assertion in this file.
"""

from collections import Counter


def test_trim_names_are_not_mangled_by_footnote_stripping(books):
    """Regression: "AT4" once became "AT" because the trailing digit was read as
    a footnote reference.  On feature sheets the digits ride on the availability
    token, so a header digit is part of the name."""
    for label in ("2026", "2027"):
        names = [c.display_name
                 for c in books[label].sheets["Standard Equipment"].columns]
        assert names == ["Elevation", "AT4", "AT4 Ultimate",
                         "Denali", "Denali Ultimate"]
        # Color and Trim is the opposite case and must still be stripped.
        colours = [c.display_name
                   for c in books[label].sheets["Color and Trim"].columns]
        assert "Dark Atmosphere / Gideon" in colours


def test_event_counts(events):
    kinds = Counter(e.kind for e in events)
    assert dict(kinds) == {
        "constraint": 40,
        "availability": 8,
        "identity": 4,
        "description": 1,
        "row_removed": 15,
        "row_added": 5,
        "membership": 9,
    }


def test_every_clause_is_typed(events):
    """Clause-level parsing leaves no untyped condition, so no judgment call is
    ever needed merely to work out *what kind* of condition something is."""
    for event in events:
        for delta in event.deltas:
            c = delta.constraints
            assert not [x for x in c.added + c.removed
                        if x.relation == "unstructured"]


def test_rewordings_are_separated_from_real_constraint_changes(events):
    """A reworded clause has the same relation and cites the same codes, so
    plain set-difference would report it as a removal plus an addition and
    invent a constraint change.  Those are the cases -- and the only cases --
    that need a judgment call."""
    added = removed = rewritten = 0
    for event in events:
        for delta in event.deltas:
            added += len(delta.constraints.added)
            removed += len(delta.constraints.removed)
            rewritten += len(delta.constraints.rewritten)
    assert (added, removed, rewritten) == (120, 179, 39)

    # Only rewrites of order-restricting relations are worth asking about;
    # reworded boilerplate cannot change what is orderable.
    assert sum(len(e.needs_constraint_judgment) for e in events) == 25

    wording_only = [
        e for e in events
        if e.kind == "constraint"
        and all(d.constraints.wording_only for d in e.deltas if d.changed)
    ]
    assert {e.code for e in wording_only} == {"B58", "NC7"}


def test_constraint_deltas_are_deterministic_across_processes(books, alignment):
    """The judgment cache is keyed on the serialized delta, so an unstable
    ordering -- or an arbitrary pairing of look-alike clauses -- would re-bill
    every judgment on each run.  Clause matching is therefore sorted and
    similarity-gated, not first-found over a set."""
    import json
    import subprocess
    import sys
    from pathlib import Path

    script = (
        "import sys, json; sys.path.insert(0, 'src');"
        "from jev_diff.parse.registry import parse_workbook;"
        "from jev_diff.align.pairing import align_books;"
        "from jev_diff.classify import build_events, consolidate;"
        "from jev_diff.judge.questions import event_state;"
        "A=parse_workbook('docs/content/GMC/2026 GMC Yukon _ Denali Export.xlsx','2026');"
        "B=parse_workbook('docs/content/GMC/2027 GMC Yukon Export.xlsx','2027');"
        "ev=consolidate(build_events(A,B,align_books(A,B)));"
        "print(json.dumps([event_state(e) for e in ev], sort_keys=True, default=str))"
    )
    root = Path(__file__).resolve().parents[1]
    runs = {
        subprocess.run([sys.executable, "-c", script], cwd=root,
                       capture_output=True, text=True, check=True).stdout
        for _ in range(3)
    }
    assert len(runs) == 1, "event state must serialize identically every run"


def test_5jl_reports_three_trims_and_leaves_denali_alone(by_code):
    """The headline case.  A raw cell diff says five changes; a strip-the-digits
    heuristic says none.  The truth is three, and the tokens never moved."""
    (event,) = by_code["5JL"]
    assert event.kind == "constraint"
    assert event.changed_trims == ["Elevation", "AT4", "AT4 Ultimate"]
    assert "Requires 20" in event.detail and "dropped" in event.detail

    for delta in event.deltas:
        assert delta.token_a == delta.token_b == "A/D", "no token ever changed"
    unchanged = {d.trim for d in event.deltas if not d.changed}
    assert unchanged == {"Denali", "Denali Ultimate"}

    dropped = [d for d in event.deltas if d.changed]
    assert all(d.constraints.loosened for d in dropped)


def test_sfe_reports_one_gained_exclusion(by_code):
    """The inverse case: one cell moves from ■3 to ■2, which looks like
    renumbering, but Denali Ultimate actually gained an exclusion."""
    (event,) = by_code["SFE"]
    assert event.kind == "constraint"
    assert event.changed_trims == ["Denali Ultimate"]
    assert "SPZ" in event.detail

    delta = next(d for d in event.deltas if d.changed)
    assert delta.constraints.tightened
    assert [c.relation for c in delta.constraints.added] == ["not_available_with"]
    assert [sorted(c.referenced_codes) for c in delta.constraints.added] == [["SPZ"]]

    # The other package trim cites a different footnote number with identical
    # text, so it must read as unchanged.
    other = next(d for d in event.deltas if d.trim == "AT4 Ultimate")
    assert not other.changed


def test_ulm_is_a_package_shift_on_denali_only(by_code):
    (event,) = by_code["ULM"]
    assert event.kind == "availability"
    assert event.changed_trims == ["Denali"]
    delta = next(d for d in event.deltas if d.changed)
    assert (delta.token_a, delta.token_b) == ("A", "■")
    assert delta.direction == "package_shift"


def test_new_prefix_only_changes_produce_no_event(by_code):
    assert "RFO" not in by_code
    assert "SFU" not in by_code


def test_identical_change_across_sheets_is_one_finding(by_code):
    """Most codes are listed on several sheets, so an edit surfaces repeatedly.
    It must count once, while still naming every sheet it touches."""
    (five_jl,) = by_code["5JL"]
    assert five_jl.sheet == "Equipment Groups"
    assert five_jl.also_on == ("Mechanical",)
    (ulm,) = by_code["ULM"]
    assert ulm.also_on == ("Interior",)


def test_orderable_status_change_is_an_identity_event(events):
    identity = [e for e in events if e.kind == "identity" and e.code]
    codes = {e.code for e in identity}
    assert codes == {"D07", "DCH", "UKL", "UW9"}
    assert all(e.order_affecting for e in identity)
    assert all("orderable" in e.detail for e in identity)


def test_spec_row_removal_and_paint_changes(events):
    specs = [e for e in events if e.sheet == "Specs"]
    assert len(specs) == 1
    assert specs[0].kind == "row_removed"
    assert '18"' in specs[0].description

    paints = {e.kind: e.description for e in events if e.sheet == "Color and Trim"}
    assert paints == {"row_removed": "Midnight Pine",
                      "row_added": "NEW! Coastal Dune"}


def test_sheets_with_no_changes_produce_no_events(events):
    touched = {e.sheet for e in events} | {s for e in events for s in e.also_on}
    for quiet in ("Dimensions", "Engine Axles", "Trailering Specs"):
        assert quiet not in touched, f"{quiet} should be unchanged"


def test_transition_lattice_directions():
    from jev_diff.classify import _direction
    from jev_diff.model.canonical import Token

    assert _direction(Token.AVAILABLE, Token.STANDARD) == "upgraded"
    assert _direction(Token.STANDARD, Token.AVAILABLE) == "decontented"
    assert _direction(Token.AVAILABLE, Token.INCLUDED) == "package_shift"
    assert _direction(Token.INCLUDED, Token.STANDARD) == "upgraded"
    assert _direction(Token.NOT_AVAILABLE, Token.STANDARD) == "content_added"
    assert _direction(Token.STANDARD, Token.NOT_AVAILABLE) == "content_removed"
    assert _direction(Token.AVAILABLE, Token.AVAILABLE_ADI) == "fulfillment_shift"
    # Both mean "cannot be had", so spelling it "--" instead of leaving it
    # blank is cosmetic, not an order-affecting change.
    assert _direction(Token.BLANK, Token.NOT_AVAILABLE) == "unchanged"
    assert _direction(Token.STANDARD, Token.STANDARD) == "unchanged"


def test_annotated_workbook_is_structurally_valid(books, events, tmp_path):
    """The annotated copy must open without Excel offering to repair it.

    V1 does not preserve embedded media -- openpyxl does not read it, so the
    17 wheel photographs are dropped. That is a documented limitation, not a
    defect, and the run warns about it; the source file is never modified.
    """
    import posixpath
    import re
    import zipfile

    from conftest import BOOKS
    from jev_diff.render.xlsx_annotate import annotate, media_count

    source = str(BOOKS["2027"])
    out = tmp_path / "annotated.xlsx"
    stats = annotate(source, str(out), events)
    assert stats["cells_filled"] > 0 and stats["log_rows"] > 0

    with zipfile.ZipFile(out) as z:
        names = set(z.namelist())
        assert z.testzip() is None
        # Every relationship target resolves.
        for name in names:
            if not name.endswith(".rels"):
                continue
            base = name.rsplit("/_rels/", 1)[0] if "/_rels/" in name else ""
            for tag in re.findall(r"<Relationship\b[^>]*>", z.read(name).decode()):
                if 'TargetMode="External"' in tag:
                    continue
                target = re.search(r'Target="([^"]+)"', tag).group(1)
                full = (target.lstrip("/") if target.startswith("/")
                        else posixpath.normpath(posixpath.join(base, target)))
                assert full in names, f"{name} -> {target}"
        # Every declared content type has a part behind it.
        declared = re.findall(r'PartName="/([^"]+)"',
                              z.read("[Content_Types].xml").decode())
        assert all(part in names for part in declared)

    import openpyxl

    a = openpyxl.load_workbook(source)
    b = openpyxl.load_workbook(out)
    assert b.sheetnames[0] == "Change Log"
    assert len(b.sheetnames) == len(a.sheetnames) + 1
    assert all(a[s].max_row == b[s].max_row for s in a.sheetnames)

    # The documented V1 loss, asserted so it cannot regress silently either way.
    assert media_count(source) == 17
    assert media_count(str(out)) == 0


def test_only_the_trims_that_changed_are_highlighted(books, events, tmp_path):
    """5JL lost a condition on three trims and kept it on two. The annotated
    copy must fill exactly the three."""
    import openpyxl

    from conftest import BOOKS
    from jev_diff.render.xlsx_annotate import annotate

    out = tmp_path / "annotated.xlsx"
    annotate(str(BOOKS["2027"]), str(out), events)
    sheet = openpyxl.load_workbook(out)["Equipment Groups"]

    filled = {ref: sheet[ref].fill.start_color.rgb
              for ref in ("D141", "E141", "F141", "G141", "H141")}
    assert all(v not in (None, "00000000") for k, v in filled.items()
               if k in ("D141", "E141", "F141")), filled
    assert all(filled[k] == "00000000" for k in ("G141", "H141")), filled
