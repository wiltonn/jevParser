"""Sheet-kind registry: workbook -> canonical model.

Deliberately GM-order-guide-aware.  A generic differ cannot see the footnote and
availability-token semantics that carry all the meaning in these files, so each
sheet shape is handled explicitly and anything unrecognised raises a warning
rather than being guessed at.

Every kind yields ``Row`` objects into one model, so align/classify/render never
branch on sheet kind.
"""

from __future__ import annotations

import hashlib
from dataclasses import dataclass, replace

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from ..model.canonical import (
    AxisEntry,
    Cell,
    Legend,
    Provenance,
    Row,
    RowImage,
    Sheet,
    Token,
    Warning_,
    Workbook,
    parse_cell_token,
)
from ..rules.default_gm import DEFAULT_GM
from ..rules.model import (
    ClauseRules,
    FeatureSheet,
    FlatSheet,
    GridSheet,
    KeyedSheet,
    RuleSet,
    SheetRegistry,
    StackedSheet,
)
from .constraints import cited_codes
from .footnotes import (
    constraints_for,
    looks_like_footer,
    normalize_description,
    normalize_label,
    parse_footer,
    parse_legend,
    resolve,
    split_description,
    strip_label_refs,
)

# Which worksheets exist and how each is shaped is ruleset data
# (``rules.SheetRegistry``).  The names below are the GM defaults in their
# original shape, kept for callers that predate rulesets.

def _names(kind: type) -> tuple[str, ...]:
    return tuple(r.name for r in DEFAULT_GM.sheets.sheets if isinstance(r, kind))


#: sheet name -> (description column index, first trim column index), 0-based.
FEATURE_SHEETS: dict[str, tuple[int, int]] = {
    r.name: (r.desc_col - 1, r.first_axis_col - 1)
    for r in DEFAULT_GM.sheets.sheets if isinstance(r, FeatureSheet)
}
#: sheet name -> header row (1-based).
KEYED_SHEETS: dict[str, int] = {
    r.name: r.header_row for r in DEFAULT_GM.sheets.sheets if isinstance(r, KeyedSheet)
}
STACKED_SHEETS = _names(StackedSheet)
GRID_SHEETS = _names(GridSheet)
FLAT_SHEETS = _names(FlatSheet)

FEATURE_HEADER_ROW = 3


def kind_of(name: str, registry: SheetRegistry | None = None) -> str:
    rule = (registry or DEFAULT_GM.sheets).rule_for(name)
    return rule.kind if rule else "unknown"


@dataclass(slots=True)
class ParseContext:
    """Everything a sheet parser needs besides the worksheet itself."""

    label: str
    path: str
    warnings: list[Warning_]
    rules: RuleSet
    rule: object = None          # the SheetRule matched for this worksheet

    @property
    def clauses(self) -> ClauseRules:
        return self.rules.clauses


def _a1(row: int, col: int) -> str:
    return f"{get_column_letter(col)}{row}"


def _text(value) -> str:
    return "" if value is None else str(value)


def _full_width_merges(ws) -> dict[int, str]:
    """Rows that are merged across (nearly) the whole sheet.

    These are ambiguous by design: on feature sheets they are section headers, on
    footnote-footer sheets they are the footnote table.  Content decides which.
    """
    out: dict[int, str] = {}
    for rng in ws.merged_cells.ranges:
        spans = rng.max_col - rng.min_col + 1
        if rng.min_row == rng.max_row and spans >= max(3, ws.max_column - 1):
            out[rng.min_row] = _text(ws.cell(rng.min_row, rng.min_col).value)
    return out


def _columns(ws, header_row: int, first_col: int, *, footer_notes=None,
             legend: Legend | None = None,
             refs_on_headers: bool = False,
             clauses: ClauseRules | None = None) -> list[AxisEntry]:
    """Build the column axis.

    Headers carry the trim name and its order code together
    (``"Denali\\n5SA1"``).  The order code is the stable identity -- it is what
    gets typed into the order system -- so it becomes ``axis_id`` when present.
    """
    cols: list[AxisEntry] = []
    for col in range(first_col + 1, ws.max_column + 1):
        raw = _text(ws.cell(header_row, col).value)
        if not raw.strip():
            continue
        lines = [p.strip() for p in raw.split("\n") if p.strip()]
        name = lines[0] if lines else raw.strip()
        order_code = lines[1] if len(lines) > 1 else None
        if refs_on_headers:
            clean_name, nums = strip_label_refs(name)
        else:
            clean_name, nums = name.strip(), ()
        refs = resolve(nums, sheet_notes=footer_notes,
                       legend_notes=legend.sheet_notes if legend else None)
        cols.append(AxisEntry(
            axis_id=order_code or normalize_label(clean_name) or clean_name.lower(),
            display_name=clean_name,
            raw_header=raw,
            index=col,
            order_code=order_code,
            refs=refs,
            conditions=constraints_for(refs, clauses),
        ))
    return cols


def _parse_feature_matrix(ws, ctx: ParseContext) -> Sheet:
    rule: FeatureSheet = ctx.rule
    wb_label, warnings, clauses = ctx.label, ctx.warnings, ctx.clauses
    desc_col, first_trim, header_row = rule.desc_col, rule.first_axis_col, rule.header_row

    legend = parse_legend(
        _text(ws.cell(1, desc_col).value), _text(ws.cell(2, desc_col).value)
    )
    columns = _columns(ws, header_row, first_trim - 1, legend=legend,
                       clauses=clauses)
    merges = _full_width_merges(ws)

    rows: list[Row] = []
    section: tuple[str, ...] = ()
    seen: dict[tuple, int] = {}
    sheet_notes: dict[int, str] = {}

    for r in range(header_row + 1, ws.max_row + 1):
        if r in merges:
            body = merges[r]
            if looks_like_footer(body):
                sheet_notes.update(parse_footer(body))
            elif body.strip():
                section = (body.strip(),)
            continue

        raw_desc = _text(ws.cell(r, desc_col).value)
        orderable = _text(ws.cell(r, rule.orderable_col).value).strip() or None
        ref = _text(ws.cell(r, rule.ref_col).value).strip() or None
        if not raw_desc.strip() and not orderable and not ref:
            continue

        base, row_notes = split_description(raw_desc)
        # A row whose only content is in column A, with no availability values,
        # is an unmerged section header (GM is inconsistent about merging them).
        if orderable and not ref and not base and not any(
            _text(ws.cell(r, c.index).value).strip() for c in columns
        ):
            section = (orderable,)
            continue

        values: dict[str, Cell] = {}
        for c in columns:
            raw = _text(ws.cell(r, c.index).value)
            token, nums, code = parse_cell_token(raw)
            refs = resolve(nums, row_notes=row_notes, sheet_notes=sheet_notes,
                           legend_notes=legend.sheet_notes)
            for bad in (x for x in refs if not x.ok):
                warnings.append(Warning_(
                    kind="dangling_footnote", severity="fatal",
                    message=(f"cell {_a1(r, c.index)} cites footnote "
                             f"{bad.number}, which no namespace defines"),
                    where=f"{wb_label}!{ws.title}!{_a1(r, c.index)}",
                ))
            if token is None:
                warnings.append(Warning_(
                    kind="unparsed_token", severity="fatal",
                    message=f"cell {_a1(r, c.index)} holds {raw!r}, not a known token",
                    where=f"{wb_label}!{ws.title}!{_a1(r, c.index)}",
                ))
            values[c.axis_id] = Cell(
                token=token if token is not None else Token.CODE,
                conditions=constraints_for(refs, clauses),
                refs=refs,
                value=code,
                prov=Provenance(wb_label, ws.title, _a1(r, c.index), r,
                                c.index, raw),
            )

        key = (ws.title, orderable, ref, normalize_description(base))
        occ = seen.get(key, 0)
        seen[key] = occ + 1
        rows.append(Row(
            sheet=ws.title, section_path=section, occurrence_index=occ,
            description=normalize_description(base), display_description=base,
            values=values, orderable_rpo=orderable, ref_rpo=ref,
            footnotes=row_notes, cited_codes=cited_codes(raw_desc, clauses),
            prov=Provenance(wb_label, ws.title, _a1(r, desc_col), r, desc_col,
                            raw_desc),
        ))

    return Sheet(name=ws.title, kind="feature_matrix", rows=rows,
                 columns=columns, legend=legend, footnotes=sheet_notes)


def _parse_keyed_rows(ws, ctx: ParseContext) -> Sheet:
    """Specs / Dimensions: a label column plus drivetrain columns.

    The label here is *data* -- "Turning diameter (with 18\" wheels)" -- so a
    changed label is a change event, not a remove-plus-add.
    """
    rule: KeyedSheet = ctx.rule
    wb_label = ctx.label
    header_row, label_col = rule.header_row, rule.label_col
    merges = _full_width_merges(ws)
    footer_notes: dict[int, str] = {
        r: body for r, body in merges.items() if looks_like_footer(body)
    }
    notes: dict[int, str] = {}
    for body in footer_notes.values():
        notes.update(parse_footer(body))

    columns = _columns(ws, header_row, rule.first_axis_col - 1, footer_notes=notes,
                       refs_on_headers=True, clauses=ctx.clauses)
    rows: list[Row] = []
    section: tuple[str, ...] = ()
    seen: dict[tuple, int] = {}

    for r in range(header_row + 1, ws.max_row + 1):
        if r in merges:
            body = merges[r]
            if not looks_like_footer(body) and body.strip():
                section = (body.strip(),)
            continue
        label = _text(ws.cell(r, label_col).value).strip()
        if not label:
            continue
        vals = [_text(ws.cell(r, c.index).value).strip() for c in columns]
        if not any(vals):
            section = (label,)          # section header without a merge
            continue

        clean, nums = strip_label_refs(label)
        refs = resolve(nums, sheet_notes=notes)
        values: dict[str, Cell] = {}
        for c, raw in zip(columns, vals):
            values[c.axis_id] = Cell(
                token=Token.CODE, conditions=(), value=raw or None,
                prov=Provenance(wb_label, ws.title, _a1(r, c.index), r,
                                c.index, raw),
            )
        key = (ws.title, normalize_label(clean))
        occ = seen.get(key, 0)
        seen[key] = occ + 1
        rows.append(Row(
            sheet=ws.title, section_path=section, occurrence_index=occ,
            description=normalize_label(clean), display_description=clean,
            values=values, footnotes=notes,
            prov=Provenance(wb_label, ws.title, _a1(r, label_col), r, label_col,
                            label),
        ))
    return Sheet(name=ws.title, kind="keyed_rows", rows=rows, columns=columns,
                 footnotes=notes)


def _parse_stacked_matrix(ws, ctx: ParseContext) -> Sheet:
    """Color and Trim: several sub-tables stacked vertically.

    Each sub-table re-introduces its own header row.  Cells hold **RPO codes**
    (``H1Y``, ``EA8``), not availability tokens.  Footnotes live in a merged
    footer row and their digits ride on row labels and column headers -- and
    across these two years the footnote *definitions were permuted*, so label
    text is unstable while the paints are not.  Rows key on the colour/seat code.
    """
    rule: StackedSheet = ctx.rule
    wb_label = ctx.label
    label_col, code_col, seat_col = rule.label_col, rule.code_col, rule.seat_col
    merges = _full_width_merges(ws)
    notes: dict[int, str] = {}
    for body in merges.values():
        if looks_like_footer(body):
            notes.update(parse_footer(body))

    # A header row is one whose first cell names a row-key dimension.
    keywords = tuple(k.lower() for k in rule.header_keywords)
    header_rows = [
        r for r in range(1, ws.max_row + 1)
        if _text(ws.cell(r, label_col).value).strip().lower() in keywords
    ]
    rows: list[Row] = []
    columns: list[AxisEntry] = []
    seen: dict[tuple, int] = {}

    for i, hr in enumerate(header_rows):
        end = header_rows[i + 1] if i + 1 < len(header_rows) else ws.max_row + 1
        part = _text(ws.cell(hr, label_col).value).strip()
        cols = _columns(ws, hr, rule.first_axis_col - 1, footer_notes=notes,
                        refs_on_headers=True, clauses=ctx.clauses)
        columns.extend(cols)
        # A code that occurs once in this sub-table identifies its row, so it can
        # act as the pairing key (paint codes: G42, G4J).  A seat code like A50
        # repeats across every decor level, so it cannot.
        code_counts: dict[str, int] = {}
        for r in range(hr + 1, end):
            c = _text(ws.cell(r, code_col).value).strip()
            if c:
                code_counts[c] = code_counts.get(c, 0) + 1
        for r in range(hr + 1, end):
            if r in merges:
                continue
            label = _text(ws.cell(r, label_col).value).strip()
            code = _text(ws.cell(r, code_col).value).strip()
            if not label and not code:
                continue
            clean, nums = strip_label_refs(label)
            refs = resolve(nums, sheet_notes=notes)
            values: dict[str, Cell] = {}
            for c in cols:
                raw = _text(ws.cell(r, c.index).value).strip()
                token, _, cell_code = parse_cell_token(raw)
                values[c.axis_id] = Cell(
                    token=token if token is not None else Token.CODE,
                    conditions=(), value=cell_code or (raw or None),
                    prov=Provenance(wb_label, ws.title, _a1(r, c.index), r,
                                    c.index, raw),
                )
            # Key on the stable code (paint code / seat code), not the label.
            key = (ws.title, part, code or normalize_label(clean))
            occ = seen.get(key, 0)
            seen[key] = occ + 1
            unique = bool(code) and code_counts.get(code) == 1
            rows.append(Row(
                sheet=ws.title, section_path=(part,), occurrence_index=occ,
                description=normalize_label(clean), display_description=clean,
                values=values, footnotes=notes,
                # Only a row-unique code becomes the pairing key.
                orderable_rpo=code if unique else None,
                cited_codes=frozenset({code} if code else ()),
                extra={"code": code,
                       "seat_trim": _text(ws.cell(r, seat_col).value).strip()},
                prov=Provenance(wb_label, ws.title, _a1(r, label_col), r, label_col,
                                label),
            ))
    return Sheet(name=ws.title, kind="stacked_matrix", rows=rows,
                 columns=columns, footnotes=notes)


def _parse_grid(ws, ctx: ParseContext) -> Sheet:
    """Engine Axles / Trailering Specs: multi-level merged headers.

    The row key is model plus engine, and the model cell is merged down over
    several engine rows, so it is carried forward.
    """
    rule: GridSheet = ctx.rule
    wb_label, warnings = ctx.label, ctx.warnings
    merges = _full_width_merges(ws)
    notes: dict[int, str] = {}
    for body in merges.values():
        if looks_like_footer(body):
            notes.update(parse_footer(body))

    header_row = next(
        (r for r in range(1, min(rule.search_rows + 1, ws.max_row + 1))
         if _text(ws.cell(r, 1).value).strip().lower() == rule.header_label.lower()),
        None,
    )
    if header_row is None:
        warnings.append(Warning_(
            kind="unrecognized_shape", severity="fatal",
            message=f"{ws.title}: no {rule.header_label.title()!r} header row found",
            where=f"{wb_label}!{ws.title}",
        ))
        return Sheet(name=ws.title, kind="grid", rows=[], columns=[])

    # Compose the two header levels so columns stay distinguishable.
    columns: list[AxisEntry] = []
    for col in range(rule.first_axis_col, ws.max_column + 1):
        upper = _text(ws.cell(header_row - 1, col).value).strip()
        lower = _text(ws.cell(header_row, col).value).strip()
        label = " / ".join(p for p in (upper, lower) if p) or f"col{col}"
        columns.append(AxisEntry(axis_id=normalize_label(label),
                                 display_name=label,
                                 raw_header=label, index=col))

    rows: list[Row] = []
    seen: dict[tuple, int] = {}
    model = ""
    for r in range(header_row + 1, ws.max_row + 1):
        if r in merges:
            continue
        first = _text(ws.cell(r, 1).value).strip()
        if first:
            model = first
        vals = {c.axis_id: _text(ws.cell(r, c.index).value).strip() for c in columns}
        if not any(vals.values()):
            continue
        engine = vals.get(columns[0].axis_id, "") if columns else ""
        label = f"{model} | {engine}".strip(" |")
        key = (ws.title, normalize_label(label))
        occ = seen.get(key, 0)
        seen[key] = occ + 1
        values = {
            c.axis_id: Cell(
                token=Token.CODE, conditions=(),
                value=_text(ws.cell(r, c.index).value).strip() or None,
                prov=Provenance(wb_label, ws.title, _a1(r, c.index), r,
                                c.index, _text(ws.cell(r, c.index).value)),
            )
            for c in columns
        }
        rows.append(Row(
            sheet=ws.title, section_path=(model,), occurrence_index=occ,
            description=normalize_label(label), display_description=label,
            values=values, footnotes=notes,
            prov=Provenance(wb_label, ws.title, _a1(r, 1), r, 1, first),
        ))
    return Sheet(name=ws.title, kind="grid", rows=rows, columns=columns,
                 footnotes=notes)


def _parse_flat(ws, ctx: ParseContext) -> Sheet:
    """The All sheet: a per-year ``code -> description`` dictionary.

    This is the arbiter for add-versus-relocate: a code present in both years'
    dictionaries cannot be a new feature, whatever the matrix rows suggest.
    """
    rule: FlatSheet = ctx.rule
    wb_label = ctx.label
    rows: list[Row] = []
    seen: dict[tuple, int] = {}
    for r in range(rule.start_row, ws.max_row + 1):
        code = _text(ws.cell(r, rule.code_col).value).strip()
        desc = _text(ws.cell(r, rule.desc_col).value).strip()
        if not code:
            continue
        key = (ws.title, code)
        occ = seen.get(key, 0)
        seen[key] = occ + 1
        rows.append(Row(
            sheet=ws.title, section_path=(), occurrence_index=occ,
            description=normalize_description(desc), display_description=desc,
            values={}, orderable_rpo=code, cited_codes=cited_codes(desc, ctx.clauses),
            prov=Provenance(wb_label, ws.title, _a1(r, rule.code_col), r,
                            rule.code_col, code),
        ))
    return Sheet(name=ws.title, kind="flat_lookup", rows=rows, columns=[])


_PARSERS = {
    "feature_matrix": _parse_feature_matrix,
    "keyed_rows": _parse_keyed_rows,
    "stacked_matrix": _parse_stacked_matrix,
    "grid": _parse_grid,
    "flat_lookup": _parse_flat,
}


def _images(path: str) -> dict[str, list[RowImage]]:
    """Index embedded images per sheet, by content hash and anchor cell.

    openpyxl loads none of these (which is exactly why the annotator patches the
    zip rather than re-saving the workbook), so the relationship graph is walked
    directly: workbook -> sheet part -> drawing -> media.

    Images are identified by content hash because filenames shift between years
    while content is identical -- and because that is what makes a placeholder
    swapped in for a real photo visible.
    """
    import re
    import zipfile

    def rels(part: str) -> dict[str, str]:
        base, _, name = part.rpartition("/")
        rel_part = f"{base}/_rels/{name}.rels"
        if rel_part not in names:
            return {}
        body = z.read(rel_part).decode("utf-8", "replace")
        return dict(re.findall(r'Id="([^"]+)"[^>]*?Target="([^"]+)"', body))

    def resolve(base_part: str, target: str) -> str:
        base = base_part.rpartition("/")[0]
        while target.startswith("../"):
            target, base = target[3:], base.rpartition("/")[0]
        return f"{base}/{target}".lstrip("/")

    out: dict[str, list[RowImage]] = {}
    with zipfile.ZipFile(path) as z:
        names = set(z.namelist())
        book = z.read("xl/workbook.xml").decode("utf-8", "replace")
        book_rels = rels("xl/workbook.xml")
        for name, rid in re.findall(
            r'<sheet[^>]*name="([^"]+)"[^>]*r:id="([^"]+)"', book
        ):
            target = book_rels.get(rid)
            if not target:
                continue
            sheet_part = resolve("xl/workbook.xml", target)
            draw_target = next(
                (t for t in rels(sheet_part).values() if "drawing" in t), None
            )
            if not draw_target:
                continue
            draw_part = resolve(sheet_part, draw_target)
            if draw_part not in names:
                continue
            draw_rels = rels(draw_part)
            xml = z.read(draw_part).decode("utf-8", "replace")
            found: list[RowImage] = []
            # Anchors and picture embeds appear in document order, so pair them.
            for anchor in re.findall(
                r"<xdr:(?:one|two)CellAnchor.*?</xdr:(?:one|two)CellAnchor>",
                xml, re.S,
            ):
                col = re.search(r"<xdr:col>(\d+)</xdr:col>", anchor)
                row = re.search(r"<xdr:row>(\d+)</xdr:row>", anchor)
                embed = re.search(r'r:embed="([^"]+)"', anchor)
                if not embed:
                    continue
                media = resolve(draw_part, draw_rels.get(embed.group(1), ""))
                if media not in names:
                    continue
                data = z.read(media)
                found.append(RowImage(
                    md5=hashlib.md5(data).hexdigest(),
                    mime=media.rsplit(".", 1)[-1].lower(),
                    size=len(data),
                    # Anchors are 0-based in the drawing XML; report 1-based.
                    anchor_row=int(row.group(1)) + 1 if row else None,
                    anchor_col=int(col.group(1)) + 1 if col else None,
                    part=media,
                ))
            if found:
                out[name] = found
    return out


def parse_worksheets(worksheets, path: str, label: str,
                     rules: RuleSet | None = None) -> tuple[dict[str, Sheet], list[Warning_]]:
    """Parse worksheet-like objects with the ruleset's sheet registry.

    Anything shaped like an openpyxl worksheet (``title``, ``cell``,
    ``max_row``, ``max_column``, ``merged_cells``) is accepted, which is what
    lets another source format reuse these parsers.
    """
    rules = rules or DEFAULT_GM
    warnings: list[Warning_] = []
    sheets: dict[str, Sheet] = {}
    for ws in worksheets:
        rule = rules.sheets.rule_for(ws.title)
        if rule is None:
            policy = rules.sheets.unknown_sheet
            if policy != "skip":
                warnings.append(Warning_(
                    kind="unknown_sheet", severity=policy,
                    message=f"no parser registered for sheet {ws.title!r}",
                    where=f"{label}!{ws.title}",
                ))
            continue
        ctx = ParseContext(label=label, path=path, warnings=warnings,
                           rules=rules, rule=rule)
        sheets[ws.title] = _PARSERS[rule.kind](ws, ctx)
    return sheets, warnings


def parse_workbook(path: str, label: str, rules: RuleSet | None = None) -> Workbook:
    """Parse one order-guide workbook into the canonical model."""
    # Opened by handle: openpyxl rejects a path without an Excel extension,
    # and stored uploads are named by content hash.
    with open(path, "rb") as fh:
        wb = load_workbook(fh, data_only=True)
        sheets, warnings = parse_worksheets(wb.worksheets, path, label, rules)
        wb.close()

    for sheet_name, imgs in _images(path).items():
        if sheet_name in sheets:
            sheets[sheet_name].images = imgs

    # Propagate intra-year multi-sheet presence onto each row.
    book = Workbook(path=path, label=label, sheets=sheets, warnings=warnings)
    idx = book.code_index()
    for sheet in sheets.values():
        sheet.rows = [
            replace(row, appears_on_sheets=frozenset(
                s for c in row.codes for s in idx.get(c, ())
            ))
            for row in sheet.rows
        ]
    return book
