"""Excel order-guide exports."""

from __future__ import annotations

import zipfile

from ..model.canonical import Workbook
from ..parse.registry import parse_workbook
from ..rules.model import RuleSet


class XlsxAdapter:
    format = "xlsx"
    extensions = (".xlsx", ".xlsm")

    def sniff(self, path: str) -> bool:
        try:
            with zipfile.ZipFile(path) as z:
                return "xl/workbook.xml" in z.namelist()
        except (zipfile.BadZipFile, OSError):
            return False

    def sheet_names(self, path: str, rules: RuleSet | None = None) -> list[str]:
        from openpyxl import load_workbook

        with open(path, "rb") as fh:
            wb = load_workbook(fh, read_only=True)
            try:
                return list(wb.sheetnames)
            finally:
                wb.close()

    def parse(self, path: str, label: str, rules: RuleSet) -> Workbook:
        return parse_workbook(path, label, rules)
